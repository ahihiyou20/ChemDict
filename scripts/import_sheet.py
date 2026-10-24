"""Import the vocabulary spreadsheet into the hardcoded lexicon.

Usage:
    python3 scripts/import_sheet.py path/to/bang-tu-vung.xlsx

Source: the Google Sheet "[Hóa ESL 11A10 + 11E2] Bảng từ vựng"
(https://docs.google.com/spreadsheets/d/169QZcTti2lvf3sVGuyJwoYf9G-zxpr0iE-0e7PzYy8Q/edit).
Download it with File > Download > Microsoft Excel (.xlsx). The tab "Bảng tổng hợp" is the
vocabulary: a "Bài N: <name>" heading opens each lesson, then a header row (column B reads
"STT") and one row per term. Columns: B STT, C Vietnamese, D English, E IPA, F part of speech,
G definition (VI), H definition (EN), I illustration credit linking to a picture, J source.
The "Bài N" tabs are only consulted to correct parts of speech (see pos_for).
Outputs shared/data/lexicon.json and WebP illustrations in public/hinh/.
"""
import hashlib
import io
import json
import re
import sys
import unicodedata
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
OUT_JSON = ROOT / "shared" / "data" / "lexicon.json"
OUT_IMG = ROOT / "public" / "hinh"
SUMMARY_TAB = "Bảng tổng hợp"
MAX_IMG_WIDTH = 1200

# Pictures pasted into the summary tab float over the wrong rows (three of them sit on the
# Bài 12 rows they were copied next to), so each one is placed by its content instead.
PICTURES = {
    "09f531ca87b11e9254e26613d144a5ad": (10, "Hydrocarbon", "Mô hình quả cầu – que của phân tử methane: một nguyên tử C nối với bốn nguyên tử H."),
    "de7c92eaa0d9e53428cee83a1ece6532": (10, "Functional group", "Mô hình phân tử ethanol kèm công thức H₃C–CH₂–OH, nhóm –OH tô màu đỏ."),
    "abba09a2cf371b13b570e461801c784b": (10, "Aromatic compound", "Công thức cấu tạo vòng benzene C₆H₆ với chú thích góc liên kết 120° và electron π không định chỗ."),
    "35c470e62782ab336ea4ffef99b5a038": (10, "Organic compound", "Mô hình quả cầu – que của phân tử butane với bốn nguyên tử C và mười nguyên tử H."),
    "6b11402c7f0d129d3c88a064423f892f": (10, "Molecular formula", "Bảng so sánh công thức phân tử C₄H₁₀ và C₂H₆O với công thức cấu tạo đầy đủ và thu gọn."),
    "3eefcc54834fa9445ab880c39f6686db": (12, "Mass spectrometry", "Hình 3.5: phổ khối lượng của neon, các vạch tại m/z 20 (90,9%), 21 (0,3%) và 22 (8,8%)."),
}

# Column I links to pictures on other sites. The site draws its own apparatus diagram for each
# (app/components/LabDiagram.vue); the key is a fragment of the linked URL.
DIAGRAMS = {
    "zjyx382": "separation",
    "T2zRyiemBWtaAXemhTe0e82AjJgCN0D4Vz6xBQ8wEQTZF5fgPDvX81i_k": "distillation",
    "m33c_slides003649x487": "evaporation",
    "ThcECeqNbRdjm": "extraction",
    "Rw2p2ncZuVBpA5oCPaD61lcTPZ422lQ6akDAmoZnsUJw": "extraction",
    "TxQRvmRSOVMm2oW": "solution",
    "CrystallizationProcess": "crystallization",
    "QNlBfhszRK4y4zfuaTKvMW5T1g7YWr9DUnjoHVnuStJxCirRooIN0BVP8": "filtration",
    "Column-Chromatography": "chromatography",
}


def clean(value):
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def fold(text):
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in text if unicodedata.category(c) != "Mn")


def slugify(text):
    return re.sub(r"[^a-z0-9]+", "-", fold(text).lower()).strip("-")


def strip_tracking(url):
    parts = urlsplit(url)
    query = [(k, v) for k, v in parse_qsl(parts.query) if not k.startswith("utm_")]
    return urlunsplit(parts._replace(query=urlencode(query)))


def build_source(cell, vi_term):
    text = clean(cell.value)
    url = strip_tracking(cell.hyperlink.target) if cell.hyperlink and cell.hyperlink.target else None
    if text and re.match(r"https?://", text):
        url = strip_tracking(text)  # the typed URL keeps fragments such as #page=70
    if not text and not url:
        return None
    label = text
    # A bare URL or a cell that repeats the term says nothing; name the site instead.
    if url and (not text or re.match(r"https?://", text) or text == vi_term):
        label = urlsplit(url).hostname
    return {"label": label, "url": url}


def lesson_tab_pos(wb):
    """Part of speech per (lesson, english, vietnamese) as the "Bài N" tabs give it."""
    found = {}
    for ws in wb.worksheets:
        m = re.fullmatch(r"\s*Bài\s+(\d+)\s*", ws.title)
        if not m:
            continue
        for r in range(1, ws.max_row + 1):
            en, vi, pos = clean(ws.cell(r, 3).value), clean(ws.cell(r, 2).value), clean(ws.cell(r, 5).value)
            if en and vi and pos:
                found[(int(m.group(1)), en.lower(), fold(vi).lower())] = pos
    return found


def pos_for(summary_pos, tab_pos):
    # The summary tab labels every adjective of Bài 11 "Verb (động từ)"; the lesson tab has them right.
    if summary_pos and summary_pos.startswith("Verb") and tab_pos and tab_pos.lower().startswith("adjective"):
        return "Adjective (tính từ)"
    return summary_pos


def diagram_for(cell):
    url = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else ""
    return next((name for key, name in DIAGRAMS.items() if key in url), None)


def save_picture(data, number, en):
    pil = Image.open(io.BytesIO(data)).convert("RGB")
    if pil.width > MAX_IMG_WIDTH:
        pil = pil.resize((MAX_IMG_WIDTH, round(pil.height * MAX_IMG_WIDTH / pil.width)), Image.LANCZOS)
    name = f"bai-{number}-{slugify(en)}.webp"
    pil.save(OUT_IMG / name, "WEBP", quality=82, method=6)
    return name, pil


def main(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[SUMMARY_TAB]
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    tab_pos = lesson_tab_pos(wb)
    pictures = {}
    for img in ws._images:
        data = img._data()  # readable once
        digest = hashlib.md5(data).hexdigest()
        if digest not in PICTURES:
            sys.exit(f"Unplaced picture {digest} near row {img.anchor._from.row + 1}: add it to PICTURES")
        number, en, alt = PICTURES[digest]
        pictures[(number, en)] = (data, alt)

    lessons = []
    for r in range(1, ws.max_row + 1):
        head = clean(ws.cell(r, 2).value)
        m = re.fullmatch(r"Bài\s+(\d+)\s*:\s*(.+)", head or "")
        if m:
            lessons.append({"number": int(m.group(1)), "title": f"Bài {m.group(1)}", "name": m.group(2), "senses": []})
            continue
        vi, en = clean(ws.cell(r, 3).value), clean(ws.cell(r, 4).value)
        if head == "STT" or not lessons or not (vi or en):
            continue
        lesson = lessons[-1]
        number = lesson["number"]
        sense = {
            "id": f"b{number}-{len(lesson['senses']) + 1:02d}",
            "lesson": number,
            "row": r,
            "vi": vi,
            "en": en,
            "ipa": clean(ws.cell(r, 5).value),
            "pos": pos_for(clean(ws.cell(r, 6).value), tab_pos.get((number, (en or "").lower(), fold(vi or "").lower()))),
            "defVi": clean(ws.cell(r, 7).value),
            "defEn": clean(ws.cell(r, 8).value),
            "illustration": None,
            "source": build_source(ws.cell(r, 10), vi),
        }
        if (number, en) in pictures:
            data, alt = pictures.pop((number, en))
            name, pil = save_picture(data, number, en)
            sense["illustration"] = {"kind": "image", "src": f"/hinh/{name}", "width": pil.width, "height": pil.height, "alt": alt}
            print(f"  picture -> Bài {number} row {r}: {en}")
        elif diagram := diagram_for(ws.cell(r, 9)):
            sense["illustration"] = {"kind": "diagram", "name": diagram}
        elif clean(ws.cell(r, 9).value):
            print(f"  warning: row {r} ({en}) links a picture with no diagram; add its URL to DIAGRAMS")
        lesson["senses"].append(sense)

    if pictures:
        sys.exit(f"Pictures without a row: {sorted(pictures)}")
    lessons.sort(key=lambda l: l["number"])
    for lesson in lessons:
        print(f"{lesson['title']} ({lesson['name']}): {len(lesson['senses'])} terms")
    OUT_JSON.write_text(
        # The sheet's name and URL stay out of the site's bundle: the site never names its source sheet.
        json.dumps({"lessons": lessons}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT_JSON.relative_to(ROOT)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
